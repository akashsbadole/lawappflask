"""
Advanced Document Management with OCR
Implements OCR-based document text extraction, AI document summarization, and advanced search
"""

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    from pdf2image import convert_from_path
except ImportError:
    convert_from_path = None

try:
    import docx
except ImportError:
    docx = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

import os
import io
import json
from flask import current_app
from models import db, Document, DocumentAnalysis, Client
import logging
from datetime import datetime
import re

# Configure logging
logger = logging.getLogger(__name__)

class DocumentOCREngine:
    """OCR engine for document text extraction"""
    
    def __init__(self, app=None):
        self.app = app
        if app is not None:
            self.init_app(app)
    
    def init_app(self, app):
        self.app = app
        # Configure Tesseract path if needed
        self.tesseract_path = app.config.get('TESSERACT_PATH', None)
        if self.tesseract_path:
            pytesseract.pytesseract.tesseract_cmd = self.tesseract_path
        
        # Supported languages
        self.supported_languages = app.config.get('OCR_LANGUAGES', ['eng', 'hin', 'mar', 'tam'])
    
    def extract_text_from_document(self, document_path, file_type):
        """Extract text from various document types"""
        try:
            if file_type.lower() in ['jpg', 'jpeg', 'png', 'bmp', 'tiff']:
                return self._extract_from_image(document_path)
            elif file_type.lower() == 'pdf':
                return self._extract_from_pdf(document_path)
            elif file_type.lower() in ['doc', 'docx']:
                return self._extract_from_docx(document_path)
            elif file_type.lower() == 'txt':
                return self._extract_from_text(document_path)
            else:
                logger.warning(f"Unsupported file type: {file_type}")
                return None
        except Exception as e:
            logger.error(f"Error extracting text from {document_path}: {str(e)}")
            return None
    
    def _extract_from_image(self, image_path):
        """Extract text from image using OCR"""
        if Image is None or pytesseract is None:
            logger.error("PIL or pytesseract not available for OCR")
            return None
            
        try:
            image = Image.open(image_path)
            
            # Preprocess image for better OCR
            image = self._preprocess_image(image)
            
            # Extract text using multiple languages
            text = pytesseract.image_to_string(
                image, 
                lang='+'.join(self.supported_languages),
                config='--psm 6 --oem 3'
            )
            
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting text from image {image_path}: {str(e)}")
            return None
    
    def _extract_from_pdf(self, pdf_path):
        """Extract text from PDF file"""
        try:
            text = ""
            
            # First try to extract text directly
            if PyPDF2 is not None:
                try:
                    with open(pdf_path, 'rb') as file:
                        pdf_reader = PyPDF2.PdfReader(file)
                        for page in pdf_reader.pages:
                            text += page.extract_text() + "\n"
                except:
                    text = ""
            
            # If direct extraction fails or returns minimal text, use OCR
            if len(text.strip()) < 100 and convert_from_path is not None and pytesseract is not None:
                try:
                    images = convert_from_path(pdf_path)
                    for i, image in enumerate(images):
                        # Preprocess image
                        image = self._preprocess_image(image)
                        
                        # Extract text
                        page_text = pytesseract.image_to_string(
                            image, 
                            lang='+'.join(self.supported_languages),
                            config='--psm 6 --oem 3'
                        )
                        text += page_text + "\n"
                except Exception as e:
                    logger.error(f"Error processing PDF with OCR: {str(e)}")
            
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting text from PDF {pdf_path}: {str(e)}")
            return None
    
    def _extract_from_docx(self, docx_path):
        """Extract text from DOCX file"""
        if docx is None:
            logger.error("python-docx not available for DOCX extraction")
            return None
            
        try:
            doc = docx.Document(docx_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {docx_path}: {str(e)}")
            return None
    
    def _extract_from_text(self, text_path):
        """Extract text from plain text file"""
        try:
            with open(text_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(text_path, 'r', encoding='latin-1') as file:
                    return file.read()
            except Exception as e:
                logger.error(f"Error reading text file {text_path}: {str(e)}")
                return None
        except Exception as e:
            logger.error(f"Error reading text file {text_path}: {str(e)}")
            return None
    
    def _preprocess_image(self, image):
        """Preprocess image for better OCR accuracy"""
        try:
            # Convert to grayscale
            if image.mode != 'L':
                image = image.convert('L')
            
            # Resize if too small
            width, height = image.size
            if width < 300 or height < 300:
                scale = max(300/width, 300/height)
                new_width = int(width * scale)
                new_height = int(height * scale)
                image = image.resize((new_width, new_height), Image.Resampling.LANCZOS)
            
            return image
        except Exception as e:
            logger.error(f"Error preprocessing image: {str(e)}")
            return image

class DocumentAnalyzer:
    """AI-powered document analysis and summarization"""
    
    def __init__(self, app=None):
        self.app = app
        if app is not None:
            self.init_app(app)
    
    def init_app(self, app):
        self.app = app
        # AI service configuration
        self.ai_service = app.config.get('AI_SERVICE', 'mock')  # mock, openai, gemini
        self.api_key = app.config.get('AI_API_KEY')
    
    def analyze_document(self, document_id, text_content):
        """Analyze document content and extract insights"""
        try:
            if not text_content or len(text_content.strip()) < 10:
                return None
            
            # Perform different types of analysis
            analyses = {}
            
            # Generate summary
            summary = self._generate_summary(text_content)
            if summary:
                analyses['summary'] = summary
            
            # Extract keywords
            keywords = self._extract_keywords(text_content)
            if keywords:
                analyses['keywords'] = keywords
            
            # Extract entities (names, dates, amounts, etc.)
            entities = self._extract_entities(text_content)
            if entities:
                analyses['entities'] = entities
            
            # Detect document type
            doc_type = self._detect_document_type(text_content)
            if doc_type:
                analyses['document_type'] = doc_type
            
            # Extract legal references
            legal_refs = self._extract_legal_references(text_content)
            if legal_refs:
                analyses['legal_references'] = legal_refs
            
            # Calculate confidence score
            confidence_score = self._calculate_confidence_score(analyses)
            
            # Save analyses to database
            for analysis_type, result in analyses.items():
                document_analysis = DocumentAnalysis(
                    document_id=document_id,
                    analysis_type=analysis_type,
                    analysis_result=json.dumps(result) if isinstance(result, (dict, list)) else str(result),
                    confidence_score=confidence_score
                )
                db.session.add(document_analysis)
            
            db.session.commit()
            
            return analyses
            
        except Exception as e:
            logger.error(f"Error analyzing document {document_id}: {str(e)}")
            return None
    
    def _generate_summary(self, text, max_sentences=3):
        """Generate document summary"""
        try:
            if self.ai_service == 'openai':
                return self._generate_summary_openai(text, max_sentences)
            elif self.ai_service == 'gemini':
                return self._generate_summary_gemini(text, max_sentences)
            else:
                # Mock summary using extractive summarization
                return self._generate_summary_mock(text, max_sentences)
        except Exception as e:
            logger.error(f"Error generating summary: {str(e)}")
            return None
    
    def _generate_summary_mock(self, text, max_sentences):
        """Mock summary using extractive summarization"""
        sentences = re.split(r'[.!?]+', text)
        sentences = [s.strip() for s in sentences if len(s.strip()) > 20]
        
        if len(sentences) <= max_sentences:
            return sentences
        
        # Score sentences based on length and position
        scored_sentences = []
        for i, sentence in enumerate(sentences):
            score = len(sentence) * (1 - i / len(sentences))  # Prefer earlier sentences
            scored_sentences.append((score, sentence))
        
        # Sort by score and take top sentences
        scored_sentences.sort(reverse=True, key=lambda x: x[0])
        top_sentences = [sent[1] for sent in scored_sentences[:max_sentences]]
        
        # Reorder by original position
        summary = []
        for sentence in sentences:
            if sentence in top_sentences:
                summary.append(sentence)
                if len(summary) >= max_sentences:
                    break
        
        return summary
    
    def _generate_summary_openai(self, text, max_sentences):
        """Generate summary using OpenAI API"""
        # Placeholder for OpenAI integration
        # You would need to install openai package and configure API key
        pass
    
    def _generate_summary_gemini(self, text, max_sentences):
        """Generate summary using Google Gemini API"""
        # Placeholder for Gemini integration
        # You would need to install google-generativeai package and configure API key
        pass
    
    def _extract_keywords(self, text):
        """Extract keywords from text"""
        try:
            # Simple keyword extraction using frequency and legal terms
            legal_terms = [
                'agreement', 'contract', 'court', 'judge', 'law', 'legal', 'party',
                'plaintiff', 'defendant', 'attorney', 'lawyer', 'evidence', 'testimony',
                'judgment', 'verdict', 'appeal', 'motion', 'hearing', 'trial',
                'settlement', 'damages', 'liability', 'breach', 'clause', 'provision',
                'statute', 'regulation', 'compliance', 'jurisdiction', 'precedent'
            ]
            
            words = re.findall(r'\b\w+\b', text.lower())
            word_freq = {}
            
            for word in words:
                if len(word) > 3 and word not in ['the', 'and', 'for', 'are', 'with', 'this', 'that', 'from']:
                    word_freq[word] = word_freq.get(word, 0) + 1
            
            # Prioritize legal terms
            keywords = []
            for term in legal_terms:
                if term in word_freq:
                    keywords.append((term, word_freq[term]))
            
            # Add other frequent words
            for word, freq in sorted(word_freq.items(), key=lambda x: x[1], reverse=True)[:20]:
                if word not in [kw[0] for kw in keywords]:
                    keywords.append((word, freq))
            
            return keywords[:15]  # Return top 15 keywords
            
        except Exception as e:
            logger.error(f"Error extracting keywords: {str(e)}")
            return []
    
    def _extract_entities(self, text):
        """Extract named entities from text"""
        entities = {}
        
        try:
            # Extract dates
            date_pattern = r'\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b|\b\d{4}[/-]\d{1,2}[/-]\d{1,2}\b'
            dates = re.findall(date_pattern, text)
            if dates:
                entities['dates'] = dates
            
            # Extract monetary amounts
            amount_pattern = r'[₹$€£]\s*\d+(?:,\d{3})*(?:\.\d{2})?|\d+(?:,\d{3})*(?:\.\d{2})?\s*(?:rupees|dollars|euros|pounds|rs|USD|EUR|GBP|INR)'
            amounts = re.findall(amount_pattern, text, re.IGNORECASE)
            if amounts:
                entities['amounts'] = amounts
            
            # Extract phone numbers
            phone_pattern = r'\b(?:\+91[-\s]?|0)?[6-9]\d{9}\b|\b\d{3}[-\s]?\d{3}[-\s]?\d{4}\b'
            phones = re.findall(phone_pattern, text)
            if phones:
                entities['phone_numbers'] = phones
            
            # Extract email addresses
            email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
            emails = re.findall(email_pattern, text)
            if emails:
                entities['emails'] = emails
            
            # Extract potential names (simplified)
            name_pattern = r'\b[A-Z][a-z]+\s+[A-Z][a-z]+\b'
            names = re.findall(name_pattern, text)
            if names:
                entities['names'] = list(set(names))  # Remove duplicates
            
            return entities
            
        except Exception as e:
            logger.error(f"Error extracting entities: {str(e)}")
            return {}
    
    def _detect_document_type(self, text):
        """Detect document type based on content"""
        text_lower = text.lower()
        
        document_types = {
            'contract': ['agreement', 'contract', 'terms', 'conditions', 'party', 'signature'],
            'court_filing': ['court', 'filed', 'petition', 'motion', 'judge', 'jurisdiction'],
            'legal_notice': ['notice', 'legal', 'demand', 'breach', 'compliance'],
            'power_of_attorney': ['power of attorney', 'poa', 'attorney-in-fact', 'authorize'],
            'will': ['last will', 'testament', 'beneficiary', 'executor', 'inheritance'],
            'deed': ['deed', 'property', 'transfer', 'grantor', 'grantee'],
            'affidavit': ['affidavit', 'sworn', 'declaration', 'oath', 'notary'],
            'invoice': ['invoice', 'bill', 'amount due', 'payment terms'],
            'receipt': ['receipt', 'received', 'payment', 'acknowledged'],
            'correspondence': ['dear', 'sincerely', 'regards', 'yours truly']
        }
        
        scores = {}
        for doc_type, keywords in document_types.items():
            score = sum(1 for keyword in keywords if keyword in text_lower)
            scores[doc_type] = score
        
        if max(scores.values()) == 0:
            return 'unknown'
        
        return max(scores, key=scores.get)
    
    def _extract_legal_references(self, text):
        """Extract legal references like case numbers, statutes, etc."""
        references = {}
        
        try:
            # Extract case numbers
            case_pattern = r'\b(?:Case|Civil|Criminal|Suit|Petition)\s+No\.?\s*[\d/-]+'
            case_numbers = re.findall(case_pattern, text, re.IGNORECASE)
            if case_numbers:
                references['case_numbers'] = case_numbers
            
            # Extract section references
            section_pattern = r'\bSection\s+\d+\s*(?:of\s*[A-Za-z\s]+)?'
            sections = re.findall(section_pattern, text, re.IGNORECASE)
            if sections:
                references['sections'] = sections
            
            # Extract act/statute references
            act_pattern = r'\b[A-Za-z\s]+Act\s*,?\s*\d{4}\b'
            acts = re.findall(act_pattern, text)
            if acts:
                references['acts'] = acts
            
            return references
            
        except Exception as e:
            logger.error(f"Error extracting legal references: {str(e)}")
            return {}
    
    def _calculate_confidence_score(self, analyses):
        """Calculate confidence score for the analysis"""
        if not analyses:
            return 0.0
        
        # Base score on number of successful analyses
        score = min(len(analyses) * 0.2, 1.0)
        
        # Boost score if we have a good summary
        if 'summary' in analyses and analyses['summary']:
            score += 0.2
        
        # Boost score if we have keywords
        if 'keywords' in analyses and len(analyses['keywords']) > 5:
            score += 0.1
        
        # Boost score if we have entities
        if 'entities' in analyses and analyses['entities']:
            score += 0.1
        
        return min(score, 1.0)

class AdvancedDocumentSearch:
    """Advanced search functionality for documents"""
    
    def __init__(self, app=None):
        self.app = app
    
    def search_documents(self, query, user_id=None, filters=None):
        """Search documents with advanced filters"""
        try:
            # Base query
            query_obj = Document.query
            
            # Apply user filter if specified
            if user_id:
                query_obj = query_obj.filter(
                    (Document.uploaded_by == user_id) |
                    (Document.client_id.in_(
                        db.session.query(Client.id).filter_by(assigned_lawyer_id=user_id)
                    ))
                )
            
            # Apply text search
            if query:
                search_terms = query.split()
                search_conditions = []
                
                for term in search_terms:
                    search_conditions.append(
                        Document.title.contains(term)
                    )
                    search_conditions.append(
                        Document.description.contains(term)
                    )
                    search_conditions.append(
                        Document.ocr_text.contains(term)
                    )
                    search_conditions.append(
                        Document.tags.contains(term)
                    )
                
                if search_conditions:
                    from sqlalchemy import or_
                    query_obj = query_obj.filter(or_(*search_conditions))
            
            # Apply additional filters
            if filters:
                if filters.get('category'):
                    query_obj = query_obj.filter(Document.category == filters['category'])
                
                if filters.get('client_id'):
                    query_obj = query_obj.filter(Document.client_id == filters['client_id'])
                
                if filters.get('case_id'):
                    query_obj = query_obj.filter(Document.case_id == filters['case_id'])
                
                if filters.get('date_from'):
                    query_obj = query_obj.filter(Document.created_at >= filters['date_from'])
                
                if filters.get('date_to'):
                    query_obj = query_obj.filter(Document.created_at <= filters['date_to'])
                
                if filters.get('file_type'):
                    query_obj = query_obj.filter(Document.file_type.like(f"%{filters['file_type']}%"))
            
            # Execute query
            documents = query_obj.order_by(Document.created_at.desc()).all()
            
            # Calculate relevance scores
            scored_documents = []
            for doc in documents:
                score = self._calculate_relevance_score(doc, query, filters)
                scored_documents.append((doc, score))
            
            # Sort by relevance score
            scored_documents.sort(key=lambda x: x[1], reverse=True)
            
            return [doc[0] for doc in scored_documents]
            
        except Exception as e:
            logger.error(f"Error searching documents: {str(e)}")
            return []
    
    def _calculate_relevance_score(self, document, query, filters=None):
        """Calculate relevance score for document"""
        score = 0.0
        
        if query:
            search_terms = query.split()
            content = f"{document.title} {document.description or ''} {document.ocr_text or ''} {document.tags or ''}"
            content_lower = content.lower()
            
            # Score based on term frequency
            for term in search_terms:
                term_lower = term.lower()
                count = content_lower.count(term_lower)
                score += count * 0.1
                
                # Boost score for title matches
                if term_lower in document.title.lower():
                    score += 0.5
                
                # Boost score for tag matches
                if document.tags and term_lower in document.tags.lower():
                    score += 0.3
        
        # Boost score for recently uploaded documents
        days_since_upload = (datetime.utcnow() - document.created_at).days
        recency_boost = max(0, 1 - days_since_upload / 365)  # Decay over a year
        score += recency_boost * 0.2
        
        # Boost score for documents with analyses
        if document.analyses:
            score += 0.1
        
        return score

# Initialize OCR and analysis engines
ocr_engine = DocumentOCREngine()
document_analyzer = DocumentAnalyzer()
document_search = AdvancedDocumentSearch()

def process_document_ocr(document_id):
    """Process document OCR in background"""
    try:
        document = Document.query.get(document_id)
        if not document:
            logger.error(f"Document {document_id} not found")
            return
        
        # Extract text using OCR
        extracted_text = ocr_engine.extract_text_from_document(
            document.file_path, 
            document.file_type
        )
        
        if extracted_text:
            # Update document with extracted text
            document.ocr_text = extracted_text
            
            # Analyze document content
            analyses = document_analyzer.analyze_document(document_id, extracted_text)
            
            if analyses and 'summary' in analyses:
                # Update document with AI summary
                document.ai_summary = ' '.join(analyses['summary']) if isinstance(analyses['summary'], list) else str(analyses['summary'])
            
            db.session.commit()
            
            logger.info(f"Successfully processed OCR for document {document_id}")
        else:
            logger.warning(f"No text extracted from document {document_id}")
            
    except Exception as e:
        logger.error(f"Error processing OCR for document {document_id}: {str(e)}")
        db.session.rollback()
